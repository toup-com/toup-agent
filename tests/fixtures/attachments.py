"""Generated attachment fixtures — no binary blobs in the repo.

Everything here is built from deps the test image already installs
(reportlab, Pillow, python-docx, python-pptx, openpyxl), so a fixture can
carry a checkable token and the test can assert the token survived the
round trip rather than asserting on a byte count.
"""

from __future__ import annotations

import io
import zipfile
from typing import Optional


def text_pdf(token: str, pages: int = 1) -> bytes:
    """A PDF with a real text layer containing `token` on every page."""
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for i in range(pages):
        c.drawString(72, 720, f"{token} page {i + 1}")
        c.showPage()
    c.save()
    return buf.getvalue()


def long_text_pdf(token: str, chars: int) -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    written = 0
    line = token * 10
    while written < chars:
        c.drawString(20, 750, line[:100])
        for y in range(700, 40, -12):
            c.drawString(20, y, line[:100])
            written += 100
            if written >= chars:
                break
        c.showPage()
    c.save()
    return buf.getvalue()


def scanned_pdf(label: str = "SCANNED") -> bytes:
    """A PDF whose only content is an IMAGE of text — no text layer at all."""
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (900, 500), "white")
    ImageDraw.Draw(img).text((20, 20), label, fill="black")
    ib = io.BytesIO()
    img.save(ib, format="PNG")

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawImage(ImageReader(io.BytesIO(ib.getvalue())), 40, 400, width=450, height=250)
    c.showPage()
    c.save()
    return buf.getvalue()


def encrypted_pdf(password: str = "hunter2") -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, encrypt=password)
    c.drawString(72, 720, "SECRET")
    c.showPage()
    c.save()
    return buf.getvalue()


def png_image(w: int = 64, h: int = 48, color: str = "red") -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, format="PNG")
    return buf.getvalue()


def jpeg_image(w: int = 2000, h: int = 1500) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (w, h), "blue").save(buf, format="JPEG", quality=85)
    return buf.getvalue()


def docx_doc(token: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(token)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pptx_deck(token: str) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = token
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def xlsx_sheet(token: str) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = token
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def zip_of(name: str, payload: bytes) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(name, payload)
    return buf.getvalue()


def corrupt_pdf() -> bytes:
    return b"%PDF-1.4\n" + b"\x00\xff" * 200


def local_storage(tmp_path) -> None:
    """Point the file_storage singleton at a temp root for this test."""
    from app.services import file_storage

    file_storage._backend = file_storage.LocalDiskBackend(root=str(tmp_path))


def have(module: str) -> bool:
    try:
        __import__(module)
        return True
    except Exception:
        return False


def maybe_skip(module: str) -> Optional[str]:
    return None if have(module) else f"{module} is not installed in this environment"
