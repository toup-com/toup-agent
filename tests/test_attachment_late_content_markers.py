"""Late content in a DOCX, PPTX, XLSX or CSV reaches the model, or the gap is said.

Before this file, the only office-format coverage was a ONE-token fixture per
format (`test_attachment_ingest.py`): a file whose first line is its last line
proves nothing about the tail. The full-file analysis of those formats had a
part-count test for DOCX (`test_attachment_analysis_budget.py`) and none that
looked at WHAT reached the model, and CSV had no test at all.

Each test here builds a small synthetic file in memory with a unique marker as
the LAST thing in it (final paragraph, last slide, last row of the last sheet,
last CSV row) behind enough numbered filler to need more than one analysis
part, then drives the real code:

* chat turn — `chat_attachments.ingest_and_store` (the upload route's ingest)
  → `build_turn_record` → `AgentRunner._build_attachment_blocks`, i.e. the
  exact text block the chat model is given;
* full-file analysis — `attachment_analysis.start_analysis` → `run_analysis`
  → the real `AnalysisModel`, whose OpenAI client is replaced by one that
  records every request. Assertions are on the recorded prompts.

Where a limit truncates by design (the 200,000-character chat preview) the
test asserts the disclosure the code makes, never a full read. Empty files
(a picture-only deck, a blank CSV, a workbook with no cell text) are `empty`
in chat and `empty_document` in analysis with no model call; nothing here
claims a full analysis of a picture-only deck or of a ZIP.

Every name, marker and cell value is synthetic.

Lane: platform (CI's sweep for an untagged file); also passes with RUN_MODE
unset and RUN_MODE=agent. Job state is kept on local storage in every lane by
patching `AA._agent_mode`, so no AGENT_ONLY table is needed.
"""

from __future__ import annotations

import io
import json
import os
import re
import sys
from types import SimpleNamespace

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "fixtures"))
import attachments as fx  # noqa: E402

from app.agent import attachment_analysis as AA  # noqa: E402
from app.agent.agent_runner import AgentRunner  # noqa: E402
from app.agent.attachment_ingest import MODEL_GUIDANCE  # noqa: E402
from app.agent.attachment_limits import MAX_EXTRACTED_CHARS_PER_DOCUMENT  # noqa: E402
from app.api import chat_attachments as CA  # noqa: E402

USER = "late-marker-user-0001"
TASK = "Summarize this whole file"
FA_TASK = "این فایل را خلاصه کن"

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV = "text/csv"

_MARKER_RE = re.compile(r"LATE-MARKER-[A-Z]+-Q7Z")

UNSUPPORTED_ANALYSIS = (
    "ERROR: Full-file analysis supports PDF, DOCX, PPTX, XLSX, plain text, "
    "Markdown, CSV, and JSON.")


# ── fixtures ─────────────────────────────────────────────────────────


@pytest.fixture(autouse=True)
def _isolated(tmp_path, monkeypatch):
    fx.local_storage(tmp_path)
    AA._active.clear()
    AA._start_locks.clear()
    CA._inflight.clear()
    # Local checkpoints in every lane (the agent-DB path has its own suites),
    # and no platform budget lookup: this file is about coverage, not budget.
    monkeypatch.setattr(AA, "_agent_mode", lambda: False)
    monkeypatch.setattr(AA, "_bundle_mode", lambda: False)
    # _persist starts a LibreOffice preview conversion for DOCX/PPTX; that is
    # latency work unrelated to what the model reads.
    from app.agent import doc_generators
    monkeypatch.setattr(doc_generators, "_prewarm_preview", lambda *a, **k: None)
    yield
    from app.services import file_storage
    file_storage._backend = None


class RecordingClient:
    """Stands in for the OpenAI client under the REAL `AnalysisModel`.

    Records every request's user content. Replies name the part it was given
    and echo any marker the request carried, so a marker read in the last part
    has to travel through the section, part and overview prompts to reach the
    delivered answer.
    """

    def __init__(self) -> None:
        self.chunks: dict[int, dict] = {}
        self.chunk_calls: list[int] = []
        self.stages: list[str] = []
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=self._create))

    def with_options(self, **_kwargs):
        return self

    async def close(self) -> None:
        pass

    async def _create(self, **kwargs):
        content = kwargs["messages"][1]["content"]
        payload = json.loads(content)
        seen = sorted(set(_MARKER_RE.findall(content)))
        echo = (" Found " + ", ".join(seen) + ".") if seen else ""
        if "chunk_number" in payload:
            number = payload["chunk_number"]
            self.chunk_calls.append(number)
            self.chunks[number] = payload
            text = f"Part {number} of {payload['chunk_count']} read.{echo}"
        else:
            stage = ("section" if "units" in payload else
                     "volume" if "sections" in payload else "overview")
            self.stages.append(stage)
            text = f"{stage.capitalize()} written.{echo}"
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=text))])


@pytest.fixture
def client(monkeypatch):
    from app.services import bundle_client

    recording = RecordingClient()
    monkeypatch.setattr(bundle_client, "make_openai_client", lambda: recording)
    return recording


# ── synthetic files ──────────────────────────────────────────────────


def _line(tag: str, i: int) -> str:
    """Numbered, so every line is a checkable token and an OOXML part of
    them stays far under the ingest layer's zip-ratio guard."""
    return f"{tag} line {i:04d} synthetic filler"


def _docx(marker: str, lines: int = 500):
    from docx import Document

    tokens = [_line("DOC", i) for i in range(1, lines + 1)]
    doc = Document()
    for i in range(0, lines, 5):
        doc.add_paragraph(". ".join(tokens[i:i + 5]) + ".")
    doc.add_paragraph(f"Closing paragraph. {marker}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), tokens


def _pptx(marker: str, slides: int = 10, per_slide: int = 40):
    from pptx import Presentation
    from pptx.util import Inches

    tokens = []
    prs = Presentation()
    for s in range(1, slides + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
        lines = [_line(f"SLIDE{s:02d}", i) for i in range(1, per_slide + 1)]
        tokens.extend(lines)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(6))
        box.text_frame.text = "\n".join(lines)
        if s == slides:
            last = slide.shapes.add_textbox(Inches(0.5), Inches(6.6), Inches(9), Inches(0.5))
            last.text_frame.text = f"Final slide footnote. {marker}"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), tokens


def _xlsx(marker: str, rows_per_sheet: int = 200):
    from openpyxl import Workbook

    tokens = []
    wb = Workbook()
    for index, title in enumerate(("Early", "Late")):
        ws = wb.active if index == 0 else wb.create_sheet(title)
        ws.title = title
        ws.append(["id", "label", "note"])
        for i in range(1, rows_per_sheet + 1):
            token = _line(title.upper(), i)
            tokens.append(token)
            ws.append([i, token, "value"])
    wb["Late"].append([rows_per_sheet + 1, "closing row", marker])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), tokens


def _csv(marker: str, rows: int = 500):
    tokens = [_line("CSV", i) for i in range(1, rows + 1)]
    body = "id,label,note\n" + "".join(f"{i},{t},value\n" for i, t in enumerate(tokens, 1))
    body += f"{rows + 1},closing row,{marker}\n"
    return body.encode("utf-8"), tokens


FORMATS = {
    # id: (builder, filename, declared mime, marker)
    "docx": (_docx, "Synthetic notes.docx", DOCX, "LATE-MARKER-DOCX-Q7Z"),
    "pptx": (_pptx, "Synthetic deck.pptx", PPTX, "LATE-MARKER-PPTX-Q7Z"),
    "xlsx": (_xlsx, "Synthetic book.xlsx", XLSX, "LATE-MARKER-XLSX-Q7Z"),
    "csv": (_csv, "Synthetic rows.csv", CSV, "LATE-MARKER-CSV-Q7Z"),
}


def _skip_without(fmt: str) -> None:
    module = {"docx": "docx", "pptx": "pptx", "xlsx": "openpyxl"}.get(fmt)
    if module and not fx.have(module):
        pytest.skip(f"{module} is not installed in this environment")


# ── helpers over the real paths ──────────────────────────────────────


async def _upload(data: bytes, name: str, mime: str) -> dict:
    """What `POST /chat/attachments` does on the agent after its gates."""
    return await CA.ingest_and_store(data, name, mime, USER)


def _chat_blocks(rec: dict) -> tuple[dict, list[str]]:
    """The turn record and the text blocks the chat model is given for it."""
    turn = CA.build_turn_record(USER, rec["attachment_id"])
    blocks = AgentRunner._build_attachment_blocks(
        AgentRunner.__new__(AgentRunner), "What does this file say?", None, records=[turn])
    return turn, [b["text"] for b in blocks if b.get("type") == "text"]


async def _analyse(rec: dict) -> dict:
    aid = rec["attachment_id"]
    state = await AA.start_analysis(USER, aid, rec, TASK)
    running = AA._active.get((USER, aid, state["analysis_id"]))
    if running is not None:
        await running
    final = AA.load_state(USER, aid, state["analysis_id"])
    assert final is not None
    return final


def _assert_every_part_read(final: dict, client: RecordingClient, name: str,
                            marker: str, tokens: list[str]) -> int:
    parts = final["page_count"]
    assert final["unit_kind"] == "chunk"
    assert final["status"] == "completed", final.get("error")
    assert parts >= 2, "the fixture must span more than one analysis part"
    # Every part sent exactly once, each told the true part count.
    assert sorted(client.chunk_calls) == list(range(1, parts + 1))
    assert {c["chunk_count"] for c in client.chunks.values()} == {parts}
    # The marker is in the LAST part's prompt, nowhere earlier.
    assert marker in client.chunks[parts]["source_text"]
    assert not any(marker in client.chunks[n]["source_text"] for n in range(1, parts))
    # Nothing between the first and the last line was dropped at a boundary.
    joined = "".join(client.chunks[n]["source_text"] for n in range(1, parts + 1))
    missing = [t for t in tokens if t not in joined]
    assert not missing, f"{len(missing)} lines never reached a prompt, e.g. {missing[:3]}"
    assert joined.index(tokens[0]) < joined.index(tokens[-1]) < joined.index(marker)
    # The synthesis saw it too, and the delivered copy claims no more than that.
    assert client.stages[-1] == "overview" and marker in final["overview"]
    pub = AA.public_state(final)
    assert (pub["unit_count"], pub["units_completed"], pub["units_failed"]) == (parts, parts, 0)
    delivered = AA._delivery_parts(final)
    assert delivered[0].startswith(f"Analysis of {name} ({parts} parts)\n\n")
    assert marker in delivered[0]
    assert "Coverage is incomplete" not in "\n".join(delivered)
    assert "not read" not in "\n".join(delivered)
    return parts


# ── late markers: chat turn ──────────────────────────────────────────


@pytest.mark.parametrize("fmt", list(FORMATS))
async def test_the_last_line_reaches_the_chat_turn(fmt):
    _skip_without(fmt)
    build, name, mime, marker = FORMATS[fmt]
    data, tokens = build(marker)
    rec = await _upload(data, name, mime)

    assert rec["mime"] == mime and rec["kind"] == "document"
    assert rec["ingest"]["status"] == "ok" and not rec["ingest"].get("truncated")
    turn, texts = _chat_blocks(rec)
    assert turn["status"] == "ok" and not turn["truncated"]
    assert rec["ingest"]["chars"] == len(turn["text"])
    body = [t for t in texts if t.startswith(f"[1] {name} — document")]
    assert len(body) == 1, texts[:2]
    body = body[0]
    assert "TRUNCATED" not in body and "COULD NOT BE READ" not in body
    assert body.index(tokens[0]) < body.index(tokens[-1]) < body.index(marker)
    assert all(t in body for t in tokens)
    assert texts[-1] == "What does this file say?"


# ── late markers: full-file analysis ─────────────────────────────────


@pytest.mark.parametrize("fmt", list(FORMATS))
async def test_the_last_line_reaches_the_last_analysis_part(fmt, client):
    _skip_without(fmt)
    build, name, mime, marker = FORMATS[fmt]
    data, tokens = build(marker)
    rec = await _upload(data, name, mime)

    final = await _analyse(rec)
    _assert_every_part_read(final, client, name, marker, tokens)


# ── a preview cut by design: said, and the analysis reads past it ────


def _long_csv(marker: str):
    rows = MAX_EXTRACTED_CHARS_PER_DOCUMENT // 30 + 400
    return _csv(marker, rows=rows)


def _long_xlsx(marker: str):
    rows = MAX_EXTRACTED_CHARS_PER_DOCUMENT // 60 + 400
    return _xlsx(marker, rows_per_sheet=rows)


@pytest.mark.parametrize("fmt,build", [("csv", _long_csv), ("xlsx", _long_xlsx)],
                         ids=["csv", "xlsx"])
async def test_a_preview_past_the_turn_cap_says_so_and_analysis_reads_the_tail(
        fmt, build, client):
    """The 200,000-character chat preview is a product limit, unchanged here.
    The chat copy must say it is cut and point at the full analysis, and that
    analysis must reach the marker the preview could not."""
    _skip_without(fmt)
    _, name, mime, marker = FORMATS[fmt]
    data, tokens = build(marker)
    rec = await _upload(data, name, mime)

    assert rec["ingest"]["status"] == "truncated" and rec["ingest"]["truncated"] is True
    turn, texts = _chat_blocks(rec)
    assert len(turn["text"]) == MAX_EXTRACTED_CHARS_PER_DOCUMENT
    body = next(t for t in texts if t.startswith(f"[1] {name} — document"))
    assert f"[1] {name} — document — attachment_id {rec['attachment_id']}, TRUNCATED" in body
    assert "This is only the chat preview." in body
    assert "call analyze_attachment with this attachment_id" in body
    assert marker not in body and tokens[-1] not in body

    final = await _analyse(rec)
    parts = _assert_every_part_read(final, client, name, marker, tokens)
    assert parts > MAX_EXTRACTED_CHARS_PER_DOCUMENT // AA.TEXT_CHUNK_CHARS


# ── empty and unsupported ────────────────────────────────────────────


def _picture_only_deck() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for _ in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(fx.png_image()), 0, 0)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _assert_no_text_told(texts: list[str], name: str) -> None:
    body = next(t for t in texts if t.startswith(f"[1] {name} — document"))
    assert "COULD NOT BE READ (empty)" in body
    assert body.endswith(MODEL_GUIDANCE["empty"])
    assert "COULD NOT BE READ (empty)" in texts[0]


async def _assert_empty_analysis(rec: dict, client: RecordingClient, name: str) -> None:
    final = await _analyse(rec)
    assert final["status"] == "failed"
    assert final["error"] == {"code": "empty_document",
                              "message": "The document contains no readable text."}
    assert final["page_count"] == 0 and final["pages"] == []
    assert client.chunk_calls == [] and client.stages == []   # no model call at all
    # The whole copy. A file with no parts at all has no "failed parts" to
    # disclaim and no numbered details to promise, in English or Persian,
    # even when details were asked for.
    body = "The document contains no readable text."
    for state in (final, dict(final, include_unit_details=True)):
        assert AA._delivery_parts(state) == [f"{name} — not read\n\n{body}"]
        assert AA._delivery_parts(dict(state, task=FA_TASK)) == [f"فایل «{name}» خوانده نشد\n\n{body}"]


@pytest.mark.skipif(not fx.have("pptx"), reason="python-pptx not installed here")
async def test_a_picture_only_deck_is_empty_in_chat_and_never_claimed_read(client):
    """Slides of pictures carry no text this path can read. The chat model is
    told the file held no readable text, and a full analysis is refused as
    empty with no model call — no slide is ever described as read."""
    name = "Pictures only.pptx"
    rec = await _upload(_picture_only_deck(), name, PPTX)
    assert rec["ingest"] == {"status": "empty", "reason": "empty", "kind": "document"}
    _, texts = _chat_blocks(rec)
    _assert_no_text_told(texts, name)
    await _assert_empty_analysis(rec, client, name)


@pytest.mark.parametrize("data", [b"", b"\r\n  \r\n"], ids=["zero-bytes", "blank-lines"])
async def test_an_empty_csv_is_empty_in_chat_and_in_analysis(data, client):
    name = "Nothing here.csv"
    rec = await _upload(data, name, CSV)
    assert rec["ingest"]["status"] == "empty"
    _, texts = _chat_blocks(rec)
    _assert_no_text_told(texts, name)
    await _assert_empty_analysis(rec, client, name)


def _workbook(sheets: dict) -> bytes:
    """{title: {(row, col): value}}; unlisted cells are blank."""
    from openpyxl import Workbook

    wb = Workbook()
    for index, (title, cells) in enumerate(sheets.items()):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = title
        for (row, col), value in cells.items():
            ws.cell(row=row, column=col, value=value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.skipif(not fx.have("openpyxl"), reason="openpyxl not installed here")
@pytest.mark.parametrize("sheets", [
    {"Sheet": {}},
    # Spaces, a tab, a no-break space and a newline, behind blank rows and
    # across two sheets: whitespace is blank, so this has no text either.
    {"Blank one": {(1, 1): "   ", (3, 2): "\t", (4, 1): " ", (4, 3): "\n"},
     "Blank two": {(2, 1): " "}},
], ids=["no-cells", "whitespace-only"])
async def test_an_empty_workbook_is_empty_in_chat_and_in_analysis(sheets, client):
    """A workbook with no cell text is `empty`, as a picture-only deck is:
    the chat model is told it held no readable text (never a sheet name
    dressed up as content) and a full analysis fails `empty_document` with no
    model call. Its sheet labels alone are not a document."""
    name = "Empty book.xlsx"
    rec = await _upload(_workbook(sheets), name, XLSX)
    assert rec["ingest"] == {"status": "empty", "reason": "empty", "kind": "document"}
    turn, texts = _chat_blocks(rec)
    assert turn["status"] == "empty" and not turn.get("text")
    _assert_no_text_told(texts, name)
    assert not any("--- Sheet:" in t for t in texts)
    await _assert_empty_analysis(rec, client, name)


@pytest.mark.skipif(not fx.have("openpyxl"), reason="openpyxl not installed here")
async def test_a_workbook_labels_only_its_sheets_with_text(client):
    """Empty sheets among sheets with text are left out of the chat preview
    and of the analysis; every sheet with text reads exactly as before."""
    name = "Mixed book.xlsx"
    rec = await _upload(_workbook({
        "Empty first": {},
        "Numbers": {(2, 1): "id", (2, 2): "amount", (3, 1): 1, (3, 2): 2.5},
        "Spaces only": {(1, 1): "  ", (5, 2): "\t"},
        "Words": {(1, 1): "alpha", (1, 3): "gamma", (4, 2): "LATE-MARKER-XLSX-Q7Z"},
        "Empty last": {},
    }), name, XLSX)
    expected = ("--- Sheet: Numbers ---\nid | amount\n1 | 2.5\n"
                "--- Sheet: Words ---\nalpha |  | gamma\n | LATE-MARKER-XLSX-Q7Z | ")
    assert rec["ingest"]["status"] == "ok"
    turn, _ = _chat_blocks(rec)
    assert turn["text"] == expected.strip()

    final = await _analyse(rec)
    assert final["status"] == "completed" and final["page_count"] == 1
    assert client.chunks[1]["source_text"] == expected


async def test_a_zip_is_previewed_in_chat_but_refused_for_full_analysis(tmp_path, client):
    """A ZIP's text entries are read into the chat preview under the ingest
    layer's archive budget, and that is all: the full-file analysis refuses
    the archive with the exact tool message, starts no job and makes no
    model call. Nothing claims the archive was analyzed in full."""
    from app.agent.tool_executor import ToolExecutor

    marker = "LATE-MARKER-ZIP-Q7Z"
    inner = "a,b\n" + "".join(f"{i},{_line('ZIP', i)}\n" for i in range(1, 51)) + f"51,{marker}\n"
    name = "Bundle.zip"
    rec = await _upload(fx.zip_of("notes.csv", inner.encode()), name, "application/zip")
    assert rec["mime"] == "application/zip" and rec["ingest"]["status"] == "ok"
    _, texts = _chat_blocks(rec)
    body = next(t for t in texts if t.startswith(f"[1] {name} — document"))
    assert "=== notes.csv ===" in body and marker in body
    assert "This is only the chat preview." not in body

    executor = ToolExecutor(workspace=str(tmp_path))
    executor.set_user_id(USER)
    result = await executor.execute("analyze_attachment", {
        "attachment_id": rec["attachment_id"], "task": TASK})
    assert result == UNSUPPORTED_ANALYSIS
    with pytest.raises(ValueError, match="^unsupported_document$"):
        await AA.start_analysis(USER, rec["attachment_id"], rec, TASK)
    assert AA._active == {}
    assert AA.load_state(USER, rec["attachment_id"], AA.analysis_id_for(TASK)) is None
    assert client.chunk_calls == [] and client.stages == []


async def test_a_zip_of_pictures_is_empty_and_refused_for_analysis(tmp_path, client):
    """Images inside an archive are not read by the chat preview (only text,
    PDF, DOCX and PPTX entries are), so a ZIP of pictures is `empty` and the
    model is told there was no readable text; full analysis refuses it."""
    from app.agent.tool_executor import ToolExecutor

    name = "Photos.zip"
    rec = await _upload(fx.zip_of("photo.png", fx.png_image()), name, "application/zip")
    assert rec["ingest"]["status"] == "empty"
    _, texts = _chat_blocks(rec)
    _assert_no_text_told(texts, name)

    executor = ToolExecutor(workspace=str(tmp_path))
    executor.set_user_id(USER)
    assert await executor.execute("analyze_attachment", {
        "attachment_id": rec["attachment_id"], "task": TASK}) == UNSUPPORTED_ANALYSIS
    assert client.chunk_calls == [] and client.stages == []
